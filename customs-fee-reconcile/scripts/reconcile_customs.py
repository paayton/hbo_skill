# -*- coding: utf-8 -*-
"""报关费对账：账单 + rows/jobs/fees/groups -> 对账结果 xlsx。

判定（用户定版 2026-10-20，21:00 修订；加拼单规则 17:29 修订）：
  - 账单侧金额 = 合计(CNY) 列（报关费+联单费+查验费+加班费；HBO 的清关费/报关费
    金额实测与账单合计构成一致，如 80+150联单=230）
  - 系统侧 = SHAYUNSHG(clientid 1332541068) 名下 feename_c='清关费/报关费'、
    未作废、未销账(paypass 空) 的应付费用；**不过滤 rpdate**（工作单费用跨月累积，
    靠子集和判定吸收，不用 cutoff）
  - **子集和判定（用户定版，同万海口径）**：拼单组账单合计在系统费用里存在
    （单笔等值或几笔之和等于账单合计）→ 一致；系统多记的其他费用不构成异常
    （实测：一工作单费用 320+140，组账单 320 → 一致）
  - 匹配键 = 提运单号（账单"提运单号"列 → HBO sono）。**不用箱号/合同号/业务编号查**
    （用户定版 2026-10-20）；HBO 无报关单号数据，报关单号仅展示
  - 拼单组（加拼单号）归组（用户 2026-10-20 17:29 定版）：
    a. **万海（WA/WI 前缀）**：万海网页版子单查询归组（见 wanhai_groups.py）
    b. **其他船司（177=MSC / COSU 等）只用系统查**：尾字母单号（如 177GPUPUS43996A）
       归到去尾字母的基数号（177GPUPUS43996）的组，工作号与基数号一致，金额按组内
       多票合计核（实测：43996+43996A 合计 160 = 系统 GOSZX26012824 费用 160）；
       COSU 纯数字超 10 位的按前 10 位数字核归组（646355248001 → 6463552480）
  - **拼单行不再特殊高亮（用户 2026-10-21 12:08 定版）**：只保留异常行黄色；归组仅用于
    正确归并/核对，加拼单号的一致行与普通一致行外观相同

输出（用户定版）：
  - 原账单不动；**账单自带 工作号/异常金额/操作 模板列时写进原列，绝不重复插入**
    （缺失的列才追加到末尾）
  - **同工作号的行归并排列**（含加拼母子单，如 177GPUPUS43996A 挪到 43996 相邻行），
    组间按最早原行号、组内保持原序，序号列保留原值；**工作号列相同值合并单元格**
    （用户 2026-10-21 09:33 定版）
  - 工作号每组都填（母单+子单同号）；异常金额只填异常行（用户 2026-10-21 15:03 定版写法）：
    「系统查无此工作单 / 系统无X（系统没有这笔票面X的费用）/ 系统Y（系统只查到Y，与票面凑不平）」
    操作(员) = 组内应付费用行的录入人英文代码（app_users.usercode，如 KATELUO/ROSIEQIN，
    去重逗号连接；组内无费用行则留空。用户 2026-10-20 18:05 定版：不用客户自编号）
  - 异常行整行黄色（FFFF00）

用法：
  python reconcile_customs.py --bill 账单.xls --rows rows.csv --jobs jobs.csv \
      --fees fees.csv --groups groups.json --out 输出.xlsx
"""
import argparse
import csv
import json
import re
from itertools import combinations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

YELLOW = PatternFill(start_color='FFFF00', fill_type='solid')
FEENAME = '清关费/报关费'


def load_maps(jobs_path, fees_path):
    jobs_by_sono = {}
    with open(jobs_path, encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            jobs_by_sono.setdefault((r['sono'] or '').strip(), []).append(
                ((r['jobno'] or '').strip(), (r['jobid'] or '').strip()))

    def find_jobs(sono):
        out = []
        for raw, js in jobs_by_sono.items():
            if raw == sono or raw.endswith(sono):
                out += js
        return sorted(set(out))

    fee_list_by_jobid = {}
    with open(fees_path, encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            if (r['feename_c'] or '').strip() != FEENAME:
                continue
            if (r.get('rptype') or '1').strip() != '1':   # 只看应付，不看应收（用户定版）
                continue
            if (r['paypass'] or '').strip():      # 已销账剔除
                continue
            cents = int(round(float(r['amt'] or 0) * 100))
            fee_list_by_jobid.setdefault((r['jobid'] or '').strip(), []).append(
                (cents, (r.get('usercode') or '').strip()))
    return find_jobs, fee_list_by_jobid


def subset_match(target_cents, amounts):
    """any subset sums to target（几笔之和=账单合计）"""
    amounts = [a for a in amounts if a > 0]
    if target_cents <= 0:
        return False
    for r in range(1, len(amounts) + 1):
        for c in combinations(amounts, r):
            if sum(c) == target_cents:
                return True
    return False


def system_group_key(sono):
    """非万海运单的系统侧归组键（用户 17:29 定版）：
    尾字母单号去尾字母（177GPUPUS43996A -> 177GPUPUS43996）；
    纯数字超 10 位取前 10 位（COSU646355248001 -> COSU6463552480）。
    返回 None = 无需归组。"""
    m = re.match(r'^([0-9A-Z]+?)(\d+)([A-Z]{1,2})$', sono)
    if m:
        return m.group(1) + m.group(2)
    m = re.match(r'^([0-9A-Z]+?)(\d{11,})$', sono)
    if m:
        return m.group(1) + m.group(2)[:10]
    return None


def build_row_groups(rows, groups_path):
    """返回 ({seq: group_key}, {group_key: set(万海组全部组员)})。
    同拼单组行共享 key；web_members 含不在本账单的组员，供费用池反查
    （组员的报关费可能记在账单外组员的工作单上，如 02月 WA666V20190 组）。"""
    parent = {r['seq']: r['seq'] for r in rows}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra

    seqs_by_sono = {}
    for r in rows:
        seqs_by_sono.setdefault(r['sono'], []).append(r['seq'])

    web_members = {}   # group_key -> set(sonos)，万海网页组全部组员（含账单外）
    # a. 万海（WA/WI）：网页版子单归组
    if groups_path:
        sono_group = {}
        with open(groups_path, encoding='utf-8') as f:
            for g in json.load(f):
                grp = g.get('group')
                if g.get('found') and grp and len(grp) > 1:
                    for so in grp:
                        sono_group[so] = (grp[0], set(grp))
        for so, (key, members) in sono_group.items():
            if so not in seqs_by_sono:
                continue  # 组里的运单不在本账单（如子单的兄弟单没进账单）
            for r in rows:
                if r['sono'] in sono_group and sono_group[r['sono']][0] == key:
                    union(r['seq'], seqs_by_sono[so][0])
        for r in rows:
            if r['sono'] in sono_group:
                key, members = sono_group[r['sono']]
                web_members.setdefault(find(r['seq']), set()).update(members)

    # b. 非万海：系统侧归组（尾字母单号归基数号 / COSU 数字核归组）
    key_by_sono = {}
    for so in seqs_by_sono:
        if so.startswith(('WA', 'WI')):
            continue  # 万海走网页，不做系统侧归组
        k = system_group_key(so)
        if k:
            key_by_sono[so] = k
    for k in {kk for kk in key_by_sono.values()}:
        # 归组成员 = 有该键的运单 + 基数号本身在账单里的行（如 177GPUPUS43996）
        members = [s for s, kk in key_by_sono.items() if kk == k]
        if k in seqs_by_sono:
            members.append(k)
        if len(members) > 1:
            for s in members[1:]:
                union(seqs_by_sono[s][0], seqs_by_sono[members[0]][0])

    # 同一运单多行（拆票申报）也归一组
    for so, seqs in seqs_by_sono.items():
        for s in seqs[1:]:
            union(seqs[0], s)
    return {r['seq']: find(r['seq']) for r in rows}, web_members


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--bill', required=True)
    ap.add_argument('--rows', required=True)
    ap.add_argument('--jobs', required=True)
    ap.add_argument('--fees', required=True)
    ap.add_argument('--groups', default=None)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()

    rows = list(csv.DictReader(open(a.rows, encoding='utf-8-sig')))
    for r in rows:
        r['total_f'] = float(r['total'] or 0) if str(r['total']).strip() else 0.0
    find_jobs, fee_list_by_jobid = load_maps(a.jobs, a.fees)
    row_group, web_members = build_row_groups(rows, a.groups)

    groups = {}
    for r in rows:
        groups.setdefault(row_group[r['seq']], []).append(r)

    result = {}
    stats = {'ok': 0, 'diff': 0, 'zero': 0, 'nojob': 0}
    for gkey, grows in groups.items():
        # 万海组：费用池覆盖全组组员（含账单外，如 WA666V20191——组员的报关费
        # 可能记在账单外组员的工作单上）；非万海组只用组内运单
        cand_sonos = set()
        for r in grows:
            cand_sonos.update(web_members.get(gkey) or {r['sono']})
        jobids, jobnos = [], []
        disp_ids = []   # 显示工作号只用账单内运单命中的工作单（账单外组员只进费用池）
        for r in grows:
            for s in [r['sono']] + ([system_group_key(r['sono'])]
                                    if system_group_key(r['sono']) not in (None, r['sono']) else []):
                for jobno, jobid in find_jobs(s):
                    if jobid not in disp_ids:
                        disp_ids.append(jobid)
        for s in sorted(cand_sonos):
            cands = [s]
            k = system_group_key(s)   # COSU 等拼单长号查不到时兜底查去尾基数号
            if k and k != s:
                cands.append(k)
            for s2 in cands:
                for jobno, jobid in find_jobs(s2):
                    if jobid not in jobids:
                        jobids.append(jobid)
                        jobnos.append(jobno)
        bill_cents = int(round(sum(r['total_f'] for r in grows) * 100))
        pool = []
        ops = []
        for j in jobids:
            for cents, ucode in fee_list_by_jobid.get(j, []):
                pool.append(cents)
                if ucode and ucode not in ops:
                    ops.append(ucode)
        operator = ','.join(ops)   # 组内应付费用行录入人去重逗号连接；无费用行则留空
        # 显示工作号：只取账单内运单命中的工作单，优先有费用的；无后缀(真单)优先
        fee_jobs = [j for j in disp_ids if fee_list_by_jobid.get(j)]
        show_ids = fee_jobs or disp_ids
        name_by_id = dict(zip(jobids, jobnos))
        disp = sorted({name_by_id[j] for j in show_ids if j in name_by_id},
                      key=lambda x: (re.search(r'[A-Z]$', x or '') is not None, x))
        workno = ','.join(disp)
        if not jobids:
            status, abnormal = 'nojob', '系统查无此工作单'
        elif subset_match(bill_cents, pool):
            status, abnormal = 'ok', ''
        elif sum(pool) == 0:
            status, abnormal = 'zero', '系统无%s' % ('%g' % (bill_cents / 100))
        else:
            status = 'diff'
            abnormal = '系统%s' % ('%g' % (sum(pool) / 100))
        stats[status] = stats.get(status, 0) + len(grows)
        is_pin = len({r['sono'] for r in grows}) > 1   # 加拼单号组（多运单归组）
        for r in grows:
            result[r['seq']] = (workno, abnormal, operator, is_pin)

    # 写输出：账单自带 工作号/异常金额/操作(员) 模板列时写进原列，缺的才追加
    src = pd.read_excel(a.bill, header=None)
    wb = Workbook()
    ws = wb.active
    hdr_row = None
    for i, row in src.iterrows():
        vals = [str(v).strip() for v in row.tolist()]
        if '提运单号' in vals and hdr_row is None:
            hdr_row = i
        for j, v in enumerate(row.tolist(), 1):
            if pd.isna(v):
                continue
            ws.cell(i + 1, j, v)
    ncol = src.shape[1]

    def find_col(names):
        for j in range(1, ncol + 1):
            v = str(ws.cell(hdr_row + 1, j).value or '').strip()
            if v in names:
                return j
        return None

    c_workno = find_col({'工作号'})
    c_abn = find_col({'异常金额'})
    c_op = find_col({'操作', '操作员'})
    if c_workno is None:
        c_workno = ncol + 1
        ws.cell(hdr_row + 1, c_workno, '工作号')
        ncol += 1
    if c_abn is None:
        c_abn = ncol + 1
        ws.cell(hdr_row + 1, c_abn, '异常金额')
        ncol += 1
    if c_op is None:
        c_op = ncol + 1
        ws.cell(hdr_row + 1, c_op, '操作员')
        ncol += 1

    # ---- 归并排序：同工作号（含加拼母子单）排在一起，工作号合并单元格 ----
    # 组间按组内最早原行号排序、组内保持原序；无工作号的行各自成组留在原位；
    # 序号列保留原值不动（便于回溯原账单），只动物理行位置
    seq_col = find_col({'序号'})
    ncol_src = src.shape[1]
    data_rows, seq_of = [], {}
    for i in range(hdr_row + 2, src.shape[0] + 1):
        seq = str(ws.cell(i, seq_col).value or '').strip()
        if seq in ('总计', '') or seq == 'None' or seq not in result:
            continue
        data_rows.append(i)
        seq_of[i] = seq
    snap = {i: [ws.cell(i, j).value for j in range(1, ncol_src + 1)]
            for i in data_rows}

    by_key = {}
    for i in data_rows:
        w = result[seq_of[i]][0]
        by_key.setdefault(w if w else ('', i), []).append(i)
    ordered = []
    for _k, idxs in sorted(by_key.items(), key=lambda kv: min(kv[1])):
        ordered += idxs

    NONE_FILL = PatternFill(fill_type=None)
    slot = hdr_row + 1
    merge_spans, run_key, run_start, run_bottom = [], None, None, None
    for i in ordered:
        slot += 1
        for j in range(1, ncol_src + 1):
            ws.cell(slot, j).value = snap[i][j - 1]
        workno, abnormal, operator, _pin = result[seq_of[i]]
        ws.cell(slot, c_workno, workno or None)
        ws.cell(slot, c_abn, abnormal or None)
        ws.cell(slot, c_op, operator or None)
        fill = YELLOW if abnormal else NONE_FILL   # 拼单不再标红（用户 2026-10-21 定版）
        for j in range(1, ncol + 1):
            ws.cell(slot, j).fill = fill
        key = workno or None
        if key and key == run_key:
            run_bottom = slot
        else:
            if run_key and run_bottom > run_start:
                merge_spans.append((run_start, run_bottom))
            if key:
                run_key, run_start, run_bottom = key, slot, slot
            else:
                run_key = None
    if run_key and run_bottom > run_start:
        merge_spans.append((run_start, run_bottom))
    for top, bottom in merge_spans:
        ws.merge_cells(start_row=top, start_column=c_workno,
                       end_row=bottom, end_column=c_workno)
    print('同工作号归并 %d 组（合并工作号单元格 %d 处）'
          % (len([1 for v in by_key.values() if len(v) > 1]), len(merge_spans)))
    wb.save(a.out)
    print('输出 %s' % a.out)
    print('一致 %d 行 / 差额 %d 行 / 系统金额无 %d 行 / 查无工作单 %d 行'
          % (stats['ok'], stats['diff'], stats['zero'], stats['nojob']))


if __name__ == '__main__':
    main()
