# -*- coding: utf-8 -*-
"""合并各月对账结果 xlsx 为一份表单：第一个 sheet「全部」= 8 个月所有账单行放在一起
（加月份列区分，工作号合并单元格逐月平移保留），其后每月一个 sheet。

用法：
  python merge_months.py --out 合并.xlsx 02月=path.xlsx 03月=path.xlsx ...
"""
import argparse
import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill(start_color='D9E1F2', fill_type='solid')
HDR_FONT = Font(bold=True)
CENTER = Alignment(horizontal='center')


def copy_cell(src, dst):
    dst.value = src.value
    if src.has_style:
        dst.fill = copy.copy(src.fill)
        dst.font = copy.copy(src.font)
        dst.alignment = copy.copy(src.alignment)
        dst.border = copy.copy(src.border)


def copy_sheet(src_ws, dst_ws):
    for row in src_ws.iter_rows():
        for c in row:
            copy_cell(c, dst_ws.cell(c.row, c.column))
    for rng in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(rng))
    for col, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[col].width = dim.width
    for r, dim in src_ws.row_dimensions.items():
        if dim.height:
            dst_ws.row_dimensions[r].height = dim.height


def append_month(all_ws, src_ws, month, row_off):
    """把 src_ws 整个 sheet（标题/表头/数据/总计）追加到 all_ws，首列插月份。
    返回新的行偏移。"""
    nrows = src_ws.max_row
    ncols = src_ws.max_column
    for i in range(1, nrows + 1):
        r = row_off + i
        mc = all_ws.cell(r, 1, month if i >= 2 else None)   # 第1行原标题行留空
        if i == 2:   # 表头行
            mc.value = '月份'
            mc.fill = copy.copy(HDR_FILL)
            mc.font = copy.copy(HDR_FONT)
            mc.alignment = copy.copy(CENTER)
        for j in range(1, ncols + 1):
            copy_cell(src_ws.cell(i, j), all_ws.cell(r, j + 1))
    for rng in src_ws.merged_cells.ranges:
        all_ws.merge_cells(start_row=rng.min_row + row_off, start_column=rng.min_col + 1,
                           end_row=rng.max_row + row_off, end_column=rng.max_col + 1)
    # 列宽（按首月为准，+1 列偏移）
    if row_off == 0:
        all_ws.column_dimensions['A'].width = 8
        for col, dim in src_ws.column_dimensions.items():
            if dim.width:
                all_ws.column_dimensions[get_column_letter(dim.index + 1)].width = dim.width
    return row_off + nrows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', required=True)
    ap.add_argument('months', nargs='+', help='月份=文件路径')
    a = ap.parse_args()

    pairs = []
    for m in a.months:
        name, _, path = m.partition('=')
        pairs.append((name, path))

    wb = Workbook()
    all_ws = wb.active
    all_ws.title = '全部'
    row_off = 0
    for name, path in pairs:
        src = load_workbook(path)
        row_off = append_month(all_ws, src.active, name, row_off)
    all_ws.freeze_panes = 'A3'

    for name, path in pairs:
        src = load_workbook(path)
        ds = wb.create_sheet(name)
        copy_sheet(src.active, ds)

    wb.save(a.out)
    total_rows = row_off
    print('合并完成 -> %s（全部 sheet %d 行 + %d 个月度 sheet）'
          % (a.out, total_rows, len(pairs)))


if __name__ == '__main__':
    main()
